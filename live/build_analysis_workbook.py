"""
Builds/updates a separate .xlsx analysis workbook from predictions_log.csv
-- log loss (model vs. de-vigged market) as real Excel formulas. This is
a SEPARATE file from predictions_readable.csv: nothing in the live/
pipeline (predict_round.py, merge_market_odds.py) ever writes to it.

APPEND-ONLY: if predictions_analysis.xlsx already exists, this only adds
rows for (round, home_team, away_team) matches not already in the sheet
-- it never rewrites, reformats, or deletes anything already there. That
means any manual formatting/colors/notes/extra columns/extra sheets you
add in Excel survive every future run untouched. The flip side: if a
match's actual_result or market odds get filled in AFTER it was already
added to the workbook, this won't go back and update that row -- edit it
by hand in Excel, or ask for a one-off targeted fix.

Log loss uses LN() formulas referencing the odds + actual_result columns
directly, so if you ever hand-correct a result in this workbook it
recalculates automatically -- same math as everywhere else in this
project: log loss = -ln(p) = ln(odds) for whichever outcome happened.
The two pooled-average summary cells (top right of the sheet) use
whole-column AVERAGE formulas, so they stay correct automatically as
rows get appended over time -- written once, on first creation, never
touched again.

Usage: run this any time predictions_log.csv changes (new round predicted,
market odds merged, results filled in) to pick up whatever's new.
"""

import os

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

_LIVE_DIR = os.path.dirname(os.path.abspath(__file__))
LOG_CSV = os.path.join(_LIVE_DIR, "predictions_log.csv")
OUTPUT_XLSX = os.path.join(_LIVE_DIR, "predictions_analysis.xlsx")
SHEET_NAME = "predictions"

COLUMNS = [
    ("round", "round"),
    ("match_date", "match_date"),
    ("home_team", "home_team"),
    ("away_team", "away_team"),
    ("odds_home", "fair_odds_home"),
    ("odds_draw", "fair_odds_draw"),
    ("odds_away", "fair_odds_away"),
    ("market_fair_odds_home", "market_fair_odds_home"),
    ("market_fair_odds_draw", "market_fair_odds_draw"),
    ("market_fair_odds_away", "market_fair_odds_away"),
    ("actual_result", "actual_result"),
]
EXTRA_HEADERS = ["log_loss_model", "log_loss_market"]

ROUND_COL, HOME_COL, AWAY_COL = 1, 3, 4  # positions of round/home_team/away_team within COLUMNS


def main():
    log = pd.read_csv(LOG_CSV)

    is_new_file = not os.path.exists(OUTPUT_XLSX)
    if is_new_file:
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET_NAME
        headers = [h for h, _ in COLUMNS] + EXTRA_HEADERS
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        existing_keys = set()
    else:
        wb = load_workbook(OUTPUT_XLSX)
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
        existing_keys = set()
        for r in range(2, ws.max_row + 1):
            rnd = ws.cell(row=r, column=ROUND_COL).value
            home = ws.cell(row=r, column=HOME_COL).value
            away = ws.cell(row=r, column=AWAY_COL).value
            if rnd is not None and home is not None:
                existing_keys.add((rnd, home, away))

    col_idx = {h: i + 1 for i, (h, _) in enumerate(COLUMNS)}
    odds_home_col = get_column_letter(col_idx["odds_home"])
    odds_draw_col = get_column_letter(col_idx["odds_draw"])
    odds_away_col = get_column_letter(col_idx["odds_away"])
    mkt_home_col = get_column_letter(col_idx["market_fair_odds_home"])
    mkt_draw_col = get_column_letter(col_idx["market_fair_odds_draw"])
    mkt_away_col = get_column_letter(col_idx["market_fair_odds_away"])
    result_col = get_column_letter(col_idx["actual_result"])
    ll_model_col = get_column_letter(len(COLUMNS) + 1)
    ll_market_col = get_column_letter(len(COLUMNS) + 2)

    appended = 0
    for _, row in log.iterrows():
        key = (row["round"], row["home_team"], row["away_team"])
        if key in existing_keys:
            continue

        r = ws.max_row + 1
        values = [row[src] if not pd.isna(row[src]) else None for _, src in COLUMNS]
        ws.append(values)

        # log loss = ln(odds assigned to whichever outcome actually happened)
        ws[f"{ll_model_col}{r}"] = (
            f'=IF({result_col}{r}="H",LN({odds_home_col}{r}),'
            f'IF({result_col}{r}="D",LN({odds_draw_col}{r}),'
            f'IF({result_col}{r}="A",LN({odds_away_col}{r}),"")))'
        )
        ws[f"{ll_market_col}{r}"] = (
            f'=IF({result_col}{r}="H",LN({mkt_home_col}{r}),'
            f'IF({result_col}{r}="D",LN({mkt_draw_col}{r}),'
            f'IF({result_col}{r}="A",LN({mkt_away_col}{r}),"")))'
        )
        appended += 1

    if is_new_file:
        # Whole-column AVERAGE so this never needs to move/resize as rows
        # get appended in future runs -- written once, left alone after.
        stats_col = get_column_letter(len(COLUMNS) + 4)
        ws[f"{stats_col}1"] = "Pooled avg log loss (model):"
        ws[f"{stats_col}1"].font = Font(bold=True)
        ws[f"{stats_col}2"] = f"=AVERAGE({ll_model_col}:{ll_model_col})"
        ws[f"{stats_col}3"] = "Pooled avg log loss (market):"
        ws[f"{stats_col}3"].font = Font(bold=True)
        ws[f"{stats_col}4"] = f"=AVERAGE({ll_market_col}:{ll_market_col})"

        for col_cells in ws.columns:
            values_in_col = [c.value for c in col_cells if c.value is not None]
            length = max((len(str(v)) for v in values_in_col), default=10)
            ws.column_dimensions[col_cells[0].column_letter].width = min(length + 2, 22)

    wb.save(OUTPUT_XLSX)
    skipped = len(log) - appended
    print(f"{'Created' if is_new_file else 'Updated'} {OUTPUT_XLSX}: "
          f"appended {appended} new row(s), skipped {skipped} already present")


if __name__ == "__main__":
    main()
